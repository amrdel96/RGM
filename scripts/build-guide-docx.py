from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
source = root / 'docs' / 'PLATFORM_GUIDE_AR.md'
out = root.parent / 'RGM-Platform-Guide-AR.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8); sec.left_margin = Cm(2); sec.right_margin = Cm(2)
styles = doc.styles
styles['Normal'].font.name = 'Arial'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); styles['Normal'].font.size = Pt(10.5)
for name, size in [('Title', 24), ('Heading 1', 17), ('Heading 2', 13), ('Heading 3', 11)]:
    styles[name].font.name = 'Arial'; styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor(0,0,0)
    styles[name].font.bold = True

def rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr(); bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1'); pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

lines = source.read_text(encoding='utf-8').splitlines()
in_code = False; code = []
i = 0
while i < len(lines):
    line = lines[i]
    if line.startswith('```'):
        if in_code:
            p = doc.add_paragraph(); p.style = 'Normal'; run = p.add_run('\n'.join(code)); run.font.name = 'Consolas'; run.font.size = Pt(8.5)
            p.paragraph_format.left_indent = Cm(0.6); rtl(p); code=[]
        in_code = not in_code; i += 1; continue
    if in_code: code.append(line); i += 1; continue
    if not line.strip(): i += 1; continue
    if line.startswith('|') and i+1 < len(lines) and set(lines[i+1].replace('|','').replace('-','').replace(':','').strip()) == set():
        headers=[x.strip() for x in line.strip('|').split('|')]; i += 2; rows=[]
        while i < len(lines) and lines[i].startswith('|'):
            rows.append([x.strip() for x in lines[i].strip('|').split('|')]); i += 1
        table=doc.add_table(rows=1, cols=len(headers)); table.style='Table Grid'
        for j,h in enumerate(headers):
            cell=table.rows[0].cells[j]; cell.text=h; shade(cell,'182765')
            for p in cell.paragraphs:
                rtl(p)
                for r in p.runs: r.font.color.rgb=RGBColor(255,255,255); r.font.bold=True
        for row in rows:
            cells=table.add_row().cells
            for j,val in enumerate(row):
                cells[j].text=val
                for p in cells[j].paragraphs: rtl(p)
        continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:], style='Title'); rtl(p)
    elif line.startswith('## '):
        p=doc.add_paragraph(line[3:], style='Heading 1'); rtl(p)
    elif line.startswith('### '):
        p=doc.add_paragraph(line[4:], style='Heading 2'); rtl(p)
    elif line.startswith('- '):
        p=doc.add_paragraph(line[2:], style='List Bullet'); rtl(p)
    elif line[:2].isdigit() and '. ' in line[:5]:
        p=doc.add_paragraph(line.split('. ',1)[1], style='List Number'); rtl(p)
    else:
        p=doc.add_paragraph(line); rtl(p)
    i += 1

for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.text='RGM Platform Guide'; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.save(out)
print(out)
